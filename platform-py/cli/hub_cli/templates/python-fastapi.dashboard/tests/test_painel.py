"""Painel (kit de dashboards) sobre o app real: conexão SQL local (mesma do `hub dev`), filtro, KPIs,
gráficos SVG, tabela, exportação CSV/XLSX e o aviso quando a conexão não está vinculada."""

from __future__ import annotations

import base64
import io
import json
import sqlite3
from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from __PKG__ import painel

ROOT = Path(__file__).resolve().parents[1]
USER = base64.urlsafe_b64encode(json.dumps({"user": "ana@exemplo"}).encode()).decode().rstrip("=")
H = {"x-hub-user": USER, "x-hub-base": "/apps/__APP_ID__"}
A = "/__APP_ID__"  # caminho do app no upstream


def _app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, connections: str) -> TestClient:
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{(tmp_path / 'app.db').as_posix()}")
    monkeypatch.setenv("HUB_SCRATCH_DIR", str(tmp_path))
    monkeypatch.delenv("HUB_DB_URL", raising=False)
    hub = tmp_path / ".hub"
    hub.mkdir(exist_ok=True)
    con = sqlite3.connect(hub / "vendas.db")
    con.executescript((ROOT / "dev" / "vendas.sql").read_text(encoding="utf-8"))
    con.close()
    (hub / "connections.toml").write_text(connections, encoding="utf-8")
    monkeypatch.setenv("HUB_SQL_LOCAL_FILE", str(hub / "connections.toml"))
    painel.BOARD.clear()
    from __PKG__.app import create_app

    return TestClient(create_app())


@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[TestClient]:
    toml = (ROOT / "dev" / "connections.toml").read_text(encoding="utf-8")
    with _app(tmp_path, monkeypatch, toml) as c:
        yield c


def test_painel_kpis_charts_table_and_filter(client: TestClient) -> None:
    r = client.get(f"{A}/painel", headers=H)
    assert r.status_code == 200, r.text
    html = r.text
    assert html.count('class="hub-dash-kpi"') == 4 and "Receita" in html
    assert "hub-dash-area" in html and "hub-dash-k-donut" in html and 'role="img"' in html
    assert 'data-hub-filter="#vendas-detalhe"' in html and "data-hub-table" in html
    assert 'action="/apps/__APP_ID__/painel"' in html and 'method="get"' in html
    assert "/_frag/__APP_ID__/painel/exportar?format=csv" in html
    assert 'href="/__APP_ID__/' not in html
    assert "<script" not in html
    todas = html.count("<tr>")
    sul = client.get(f"{A}/painel", params={"regiao": "Sul"}, headers=H).text
    assert sul.count("<tr>") < todas and '<option value="Sul" selected>' in sul
    # filtro inválido cai no padrão (nunca vai cru para o SQL)
    hostil = client.get(f"{A}/painel", params={"regiao": "x' or 1=1"}, headers=H).text
    assert hostil.count("<tr>") == todas


def test_export_csv_and_xlsx(client: TestClient) -> None:
    c = client.get(f"{A}/painel/exportar", params={"regiao": "Norte"}, headers=H)
    assert c.status_code == 200 and c.headers["content-disposition"].endswith('vendas.csv"')
    lines = c.content.decode("utf-8").splitlines()
    assert lines[0].endswith("mes;regiao;produto;pedidos;valor") and all(";Norte;" in x for x in lines[1:])
    x = client.get(f"{A}/painel/exportar", params={"format": "xlsx"}, headers=H)
    assert x.headers["content-disposition"].endswith('vendas.xlsx"')
    from openpyxl import load_workbook

    rows = list(load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True))
    assert rows[0] == ("mes", "regiao", "produto", "pedidos", "valor") and len(rows) == 73


def test_painel_without_connection_explains(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    with _app(tmp_path, monkeypatch, "") as c:
        r = c.get(f"{A}/painel", headers=H)
        assert r.status_code == 200 and "Vincule a conexão vendas" in r.text and "hub-dash-kpi" not in r.text
        assert c.get(f"{A}/painel/exportar", headers=H).status_code == 404
