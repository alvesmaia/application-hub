"""hub_sdk.dashboards / hub_sdk.sql: SVG charts (snapshots), KPIs, filters, cache, export and connections
local (``.hub/connections.toml``) and through the broker (fake HTTP server)."""

from __future__ import annotations

import asyncio
import io
import json
import os
import sqlite3
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, ClassVar

import pytest
from typing_extensions import override

from hub_sdk import dashboards as dash
from hub_sdk import sql
from hub_sdk.sql import SqlResult

SNAP = Path(__file__).parent / "snapshots" / "dashboards"
R = SqlResult(
    ["mes", "norte", "sul"],
    [["2026-01", 120, 80.5], ["2026-02", 150, 95], ["2026-03", 90, 130], ["2026-04", 170, 110]],
)
PIE = SqlResult(["status", "n"], [["Ativo", 120], ["Vencido", 8], ["Em análise", 30]])


def snapshot(name: str, got: str) -> None:
    f = SNAP / f"{name}.html"
    if os.environ.get("UPDATE_SNAPSHOTS") == "1" or not f.exists():
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(got, encoding="utf-8", newline="\n")
    assert got == f.read_text(encoding="utf-8"), f"snapshot {name} mudou (UPDATE_SNAPSHOTS=1 para regravar)"


@pytest.mark.parametrize(
    ("name", "make"),
    [
        ("line", lambda: dash.line_chart(R, "mes", ["norte", "sul"], "Vendas por mês", "R$")),
        ("area", lambda: dash.area_chart(R, "mes", "norte", "Norte")),
        ("bar", lambda: dash.bar_chart(R, "mes", ["norte", "sul"], "Comparativo")),
        ("pie", lambda: dash.pie_chart(PIE, "status", "n", "Contratos")),
        ("donut", lambda: dash.donut_chart(PIE, "status", "n", "Contratos")),
        ("table", lambda: dash.table(R, [("mes", "Mês"), ("norte", "Norte")], units={"norte": "R$"})),
        (
            "kpis",
            lambda: dash.kpis(
                [
                    dash.kpi("Receita", 1234567.5, "R$", delta=0.12),
                    dash.kpi("Pedidos", 42, hint="no mês", delta=-0.05),
                ]
            ),
        ),
    ],
)
def test_snapshots(name: str, make: Any) -> None:
    out = str(make())
    snapshot(name, out)
    assert "<script" not in out and "on" + "click" not in out and "style=" not in out


def test_chart_semantics_and_edge_cases() -> None:
    s = str(dash.line_chart(R, "mes", ["norte", "sul"], "Vendas"))
    assert 'role="img"' in s and "aria-labelledby" in s and "<desc>" in s and s.count("<circle") == 8
    assert 'class="hub-dash-legend"' in s  # 2 series → legend
    assert "<script" not in str(dash.chart("bar", ["<b>"], [("x<y", [1])]))  # rótulos escapados
    assert "&lt;b&gt;" in str(dash.chart("bar", ["<b>"], [("x", [1])]))
    assert "Sem dados" in str(dash.chart("line", [], []))
    assert "Sem dados" in str(dash.chart("pie", ["a"], [("x", [0])]))
    one = str(dash.chart("donut", ["a"], [("x", [5])]))
    assert "hub-dash-hole" in one
    many = str(dash.chart("pie", list("abcdefgh"), [("x", [1, 2, 3, 4, 5, 6, 7, 8])]))
    assert "Outros" in many
    neg = str(dash.chart("bar", ["a", "b"], [("x", [-5, "7,5"])], unit="un"))
    assert "-5 un" in neg and "7,50 un" in neg
    assert dash.fmt(None) == "–" and dash.fmt("x") == "x" and dash.fmt(3, "%") == "3 %"
    assert dash.compact(2_500_000) == "2,5 mi" and dash.compact(12) == "12"


def test_params_filters_and_export_links() -> None:
    p = [
        dash.Param("regiao", "Região", options=[("", "Todas"), ("N", "Norte")]),
        dash.Param("desde", "Desde", "date"),
        dash.Param("min", "Mínimo", "number", default="0"),
        dash.Param("q", "Busca", "text"),
    ]
    board = dash.Dashboard("x", [], p)
    v = board.values({"regiao": "X'; drop", "desde": "2026-02-30", "min": "2.5", "q": "a\x00b"})
    assert v == {"regiao": "", "desde": "", "min": 2.5, "q": "ab"}
    assert board.values({"regiao": "N", "desde": "2026-02-01", "min": "3"})["min"] == 3
    f = str(dash.filters_form(p, {"regiao": "N"}, "/apps/x"))
    assert 'method="get"' in f and 'action="/apps/x"' in f and '<option value="N" selected>' in f
    assert 'type="date"' in f and 'type="number"' in f
    links = str(dash.export_links("/apps/x/exportar", {"regiao": "N", "q": ""}))
    assert "format=csv" in links and "regiao=N" in links and "q=" not in links


def test_csv_xlsx_and_response() -> None:
    hostile = SqlResult(["a", "b"], [["=1+1", None], ["ok;x", 2]])
    c = dash.to_csv(hostile).decode("utf-8")
    assert c.startswith("﻿a;b") and "'=1+1" in c and '"ok;x"' in c
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(dash.to_xlsx(hostile)))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("a", "b") and rows[1][0] == "'=1+1" and rows[2][1] == 2
    r = dash.export_response(hostile, "xlsx", "relatório")
    assert r.headers["content-disposition"].endswith('.xlsx"') and r.headers["cache-control"] == "no-store"
    assert dash.export_response(hostile, "csv", "x").headers["content-type"].startswith("text/csv")


def _sample(tmp: Path) -> Path:
    hub = tmp / ".hub"
    hub.mkdir()
    db = sqlite3.connect(hub / "sample.db")
    db.executescript("create table v(mes text, regiao text, valor real);"
                     "insert into v values('01','N',10),('01','S',5),('02','N',7),('03','S',2);")  # fmt: skip
    db.commit()
    db.close()
    (hub / "connections.toml").write_text(
        '[connections.vendas]\nurl = "sqlite:///.hub/sample.db"\nmax_rows = 3\n'
        '[connections.rw]\nurl = "sqlite:///.hub/sample.db"\nreadonly = false\n',
        encoding="utf-8",
    )
    return hub / "connections.toml"


def test_local_connections_dataset_and_cache(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("HUB_SQL_LOCAL_FILE", str(_sample(tmp_path)))
    assert set(sql.local_connections()) == {"vendas", "rw"}
    assert sql.connection("nada") is None
    conn = sql.connection("vendas")
    assert conn is not None and conn.for_user("u").name == "vendas"
    ds = dash.Dataset(
        "por_mes",
        "select mes, sum(valor) as total from v where (:regiao = '' or regiao = :regiao) group by mes order by mes",
        "vendas",
        params=("regiao",),
        ttl=60,
    )
    board = dash.Dashboard(
        "Vendas", [ds], [dash.Param("regiao", "Região", options=[("", "Todas"), ("N", "N")])]
    )

    async def go() -> None:
        r = await board.load("por_mes", {"regiao": ""})
        assert r.columns == ["mes", "total"] and r.rows == [["01", 15.0], ["02", 7.0], ["03", 2.0]]
        assert r.dicts()[0] == {"mes": "01", "total": 15.0}
        n = await board.load("por_mes", {"regiao": "N"})
        assert [x[0] for x in n.rows] == ["01", "02"]
        assert await board.load("por_mes", {"regiao": ""}) is r  # cache
        board.clear()
        assert await board.load("por_mes", {"regiao": ""}) is not r
        big = await conn.query("select * from v")
        assert len(big.rows) == 3 and big.truncated
        with pytest.raises(sql.SqlError) as e:
            await conn.query("delete from v")
        assert e.value.status == 403
        with pytest.raises(sql.SqlError) as e:
            await conn.query("select * from nao_existe")
        assert e.value.status == 502
        rw = sql.connection("rw")
        assert rw is not None
        await rw.query("insert into v values('04','N',1)")  # dev: somente leitura desligado
        with pytest.raises(KeyError):
            await board.load("nada")

    asyncio.run(go())


def test_from_engine(tmp_path: Path) -> None:
    from sqlalchemy import create_engine

    eng = create_engine(f"sqlite:///{(tmp_path / 'a.db').as_posix()}")
    with eng.begin() as c:
        c.exec_driver_sql("create table t(a int)")
        c.exec_driver_sql("insert into t values (1), (2)")
    conn = sql.from_engine(eng)
    ds = dash.Dataset("t", "select a from t order by a", conn, ttl=0)
    r = asyncio.run(dash.Dashboard("t", [ds]).load("t"))
    assert r.column("a") == [1, 2]
    with pytest.raises(sql.SqlError):
        dash.Dataset("x", "select 1", "nada-vinculado").connection()


class _Broker(BaseHTTPRequestHandler):
    seen: ClassVar[list[dict[str, Any]]] = []
    status = 200

    def do_POST(self) -> None:
        n = int(self.headers.get("content-length", "0"))
        body = json.loads(self.rfile.read(n))
        _Broker.seen.append(
            {
                "path": self.path,
                "auth": self.headers.get("authorization"),
                "user": self.headers.get("x-hub-user"),
                **body,
            }
        )
        if _Broker.status != 200:
            out = json.dumps({"error": "Conexão somente leitura"}).encode()
            self.send_response(_Broker.status)
        else:
            out = json.dumps({"columns": ["a"], "rows": [[1]], "truncated": False, "ms": 3}).encode()
            self.send_response(200)
        self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(out)))
        self.end_headers()
        self.wfile.write(out)

    @override
    def log_message(self, format: str, *args: Any) -> None:
        pass


def test_broker_connection(monkeypatch: pytest.MonkeyPatch) -> None:
    srv = ThreadingHTTPServer(("127.0.0.1", 0), _Broker)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        monkeypatch.delenv("HUB_SQL_LOCAL_FILE", raising=False)
        monkeypatch.setenv("HUB_CONNECTIONS", "docs=s3,vendas=postgres")
        monkeypatch.setenv("HUB_STORAGE_URL", f"http://127.0.0.1:{srv.server_port}")
        monkeypatch.setenv("HUB_STORAGE_TOKEN", "t" * 32)
        assert sql.connection("docs") is None  # file connection is not SQL
        c = sql.connection("vendas")
        assert c is not None and "postgres" in repr(c) and "t" * 32 not in repr(c)
        r = asyncio.run(c.for_user("dXNlcg", "sig").query("select :a", {"a": 1}, 10))
        assert r.rows == [[1]] and r.ms == 3
        got = _Broker.seen[-1]
        assert got["path"] == "/_hub/sql/vendas/query" and got["auth"] == "Bearer " + "t" * 32
        assert got["params"] == {"a": 1} and got["max_rows"] == 10 and got["user"] == "dXNlcg"
        _Broker.status = 403
        with pytest.raises(sql.SqlError) as e:
            asyncio.run(c.query("delete"))
        assert e.value.status == 403
    finally:
        _Broker.status = 200
        srv.shutdown()


def test_css_has_dark_theme_and_no_hardcoded_text_colors_without_tokens() -> None:
    assert ':root[data-theme="dark"]{' in dash.CSS and "prefers-color-scheme" in dash.CSS
    assert "var(--color-text-muted" in dash.CSS
